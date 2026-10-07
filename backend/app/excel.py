"""Excel templates for station records: download a pre-filled template, hand it out, upload it back.

Each template has a data sheet (row 5 onwards), an instructions sheet and a hidden "_meta" sheet that ties the file to
its kind and station, so a file filled for another station cannot be imported by mistake. Import is positional (by
column), so the header language of the file does not matter; choice cells accept uz / ru / en spellings.
"""
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FIRST_ROW = 5
BLANK_ROWS = 300
MAX_ROWS = 5000


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v).strip().lower().replace("ʼ", "'").replace("‘", "'").replace("’", "'").replace("`", "'"))


BOOL = {True: ["ha", "да", "yes", "1", "true", "bor", "есть"], False: ["yo'q", "нет", "no", "0", "false"]}

# column spec: key, width, kind, required, choices {canonical: [aliases]}, group
TRACKS = [
    ("park", 22, "text", True, None, None),
    ("track_no", 9, "text", True, None, None),
    ("purpose", 28, "text", True, None, None),
    ("arrow_from", 11, "text", False, None, "arrows"),
    ("arrow_to", 11, "text", False, None, "arrows"),
    ("length_total", 12, "int", True, None, "length"),
    ("usable_length", 12, "int", True, None, "length"),
    ("wagon_capacity", 12, "int", False, None, None),
    ("insulation", 14, "choice", False, {"Bor": ["bor", "есть", "yes", "full"], "Yo'q": ["yo'q", "нет", "no", "none"],
                                        "Qisman": ["qisman", "частично", "partial"]}, "devices"),
    ("contact_network", 14, "bool", False, None, "devices"),
    ("alns", 10, "bool", False, None, "devices"),
]
WORKERS = [
    ("full_name", 34, "text", True, None, None),
    ("position", 26, "text", True, None, None),
    ("gender", 11, "choice", True, {"M": ["erkak", "мужской", "муж", "м", "male", "m"], "F": ["ayol", "женский", "жен", "ж", "female", "f"]}, None),
    ("birth_date", 14, "date", True, None, None),
    ("education", 22, "choice", True, {"higher": ["oliy", "oliy ta'lim", "высшее", "higher"],
                                       "second": ["ikkinchi mutaxassislik", "второе высшее", "second degree", "second"],
                                       "special": ["o'rta maxsus", "среднее специальное", "vocational", "special"],
                                       "secondary": ["o'rta", "среднее", "secondary"]}, None),
    ("contract_type", 20, "text", False, None, None),
    ("hired_date", 14, "date", True, None, None),
    ("phone", 18, "text", False, None, None),
]
ROADS = [
    ("owner", 30, "text", True, None, None),
    ("inn", 13, "text", False, None, None),
    ("leader", 26, "text", False, None, "rleader"),
    ("email", 22, "text", False, None, "rleader"),
    ("phone", 17, "text", False, None, "rleader"),
    ("distance_km", 12, "float", False, None, None),
    ("length_total", 11, "float", True, None, "rlength"),
    ("length_t1", 9, "float", False, None, "rlength"),
    ("length_t2", 9, "float", False, None, "rlength"),
    ("contract_date", 13, "date", True, None, "contract"),
    ("contract_start", 13, "date", True, None, "contract"),
    ("contract_end", 13, "date", True, None, "contract"),
    ("contract_type", 16, "text", False, None, "contract"),
    ("has_tech_passport", 12, "bool", False, None, None),
    ("has_instruction", 12, "bool", False, None, None),
    ("junction", 22, "text", False, None, None),
    ("load_hours", 10, "float", False, None, "cargo"),
    ("unload_hours", 10, "float", False, None, "cargo"),
    ("oper2_hours", 10, "float", False, None, "cargo"),
    ("front", 10, "int", False, None, None),
    ("wagons_at_once", 11, "int", False, None, None),
    ("created_year", 10, "int", False, None, None),
    ("canceled", 11, "bool", False, None, None),
]
SPECS = {"tracks": TRACKS, "workers": WORKERS, "roads": ROADS}

T = {
    "uz": {
        "roads": "Shahobcha yo'llar", "rleader": "Rahbar", "leader": "F.I.Sh.", "rlength": "Uzunligi (m)", "contract": "Shartnoma", "cargo": "Yuk operatsiyalari muddati (soat)",
        "owner": "Shahobcha egasi", "inn": "STIR", "email": "Elektron pochta", "distance_km": "Qo'yib olish masofasi (km)",
        "length_t1": "T1", "length_t2": "T2", "contract_date": "Tuzilgan", "contract_start": "Boshlanishi", "contract_end": "Tugashi",
        "has_tech_passport": "Texnik pasport", "has_instruction": "Yo'riqnoma", "junction": "Tutashuv", "load_hours": "Yuklash",
        "unload_hours": "Tushirish", "oper2_hours": "Qo'sh operatsiya", "front": "Front (vag.)", "wagons_at_once": "Bir berishda vagon",
        "created_year": "Tashkil etilgan yil", "canceled": "Bekor qilingan",
        "tracks": "Parklar va yo'llar", "workers": "Xodimlar ro'yxati", "hint": "* — majburiy ustunlar. Ma'lumotlarni 5-qatordan boshlab kiriting. Ustunlar tartibini o'zgartirmang.",
        "arrows": "Chegaralovchi strelkalar", "length": "Uzunligi (m)", "devices": "Qurilmalar",
        "park": "Park", "track_no": "Yo'l №", "purpose": "Nomlanishi va vazifasi", "arrow_from": "dan", "arrow_to": "gacha",
        "length_total": "Umumiy", "usable_length": "Foydali", "wagon_capacity": "Vagon sig'imi", "insulation": "Izolyatsiya",
        "contact_network": "Kontakt tarmog'i", "alns": "ALNS",
        "full_name": "F.I.Sh.", "position": "Lavozim", "gender": "Jinsi", "birth_date": "Tug'ilgan sana", "education": "Ma'lumoti",
        "contract_type": "Shartnoma turi", "hired_date": "Ishga qabul qilingan", "phone": "Telefon",
        "yes": "Ha", "no": "Yo'q", "M": "Erkak", "F": "Ayol", "higher": "Oliy ta'lim", "second": "Ikkinchi mutaxassislik",
        "special": "O'rta maxsus", "secondary": "O'rta",
        "guide": "Yo'riqnoma", "g": [
            "Shablon {station} (ESR {esr}) stansiyasi uchun tayyorlangan. Boshqa stansiya uchun alohida shablon yuklab oling.",
            "Birinchi varaqni to'ldiring: har bir qator — bitta yozuv. Mavjud ma'lumotlar allaqachon kiritilgan, ularni tahrirlash mumkin.",
            "Ro'yxatli kataklarda qiymatni ro'yxatdan tanlang. Sanalar: KK.OO.YYYY.",
            "Fayl yuklanganda stansiyadagi ushbu bo'lim ma'lumotlari fayldagilar bilan to'liq almashtiriladi.",
            "Xato bo'lsa, tizim qator va ustunni ko'rsatadi va hech narsa o'zgartirmaydi.",
        ],
    },
    "ru": {
        "roads": "Подъездные пути", "rleader": "Руководитель", "leader": "Ф.И.О.", "rlength": "Длина (м)", "contract": "Договор", "cargo": "Срок грузовых операций (ч)",
        "owner": "Владелец пути", "inn": "ИНН", "email": "Эл. почта", "distance_km": "Расстояние подачи (км)",
        "length_t1": "T1", "length_t2": "T2", "contract_date": "Заключён", "contract_start": "Начало", "contract_end": "Окончание",
        "has_tech_passport": "Тех. паспорт", "has_instruction": "Инструкция", "junction": "Примыкание", "load_hours": "Погрузка",
        "unload_hours": "Выгрузка", "oper2_hours": "Сдвоенная", "front": "Фронт (ваг.)", "wagons_at_once": "Вагонов за подачу",
        "created_year": "Год создания", "canceled": "Расторгнут",
        "tracks": "Парки и пути", "workers": "Список сотрудников", "hint": "* — обязательные столбцы. Данные вводите с 5-й строки. Не меняйте порядок столбцов.",
        "arrows": "Ограничивающие стрелки", "length": "Длина (м)", "devices": "Устройства",
        "park": "Парк", "track_no": "Путь №", "purpose": "Наименование и назначение", "arrow_from": "от", "arrow_to": "до",
        "length_total": "Общая", "usable_length": "Полезная", "wagon_capacity": "Вместимость (ваг.)", "insulation": "Изоляция",
        "contact_network": "Контактная сеть", "alns": "АЛСН",
        "full_name": "Ф.И.О.", "position": "Должность", "gender": "Пол", "birth_date": "Дата рождения", "education": "Образование",
        "contract_type": "Тип договора", "hired_date": "Дата приёма", "phone": "Телефон",
        "yes": "Да", "no": "Нет", "M": "Мужской", "F": "Женский", "higher": "Высшее", "second": "Второе высшее",
        "special": "Среднее специальное", "secondary": "Среднее",
        "guide": "Инструкция", "g": [
            "Шаблон подготовлен для станции {station} (ЕСР {esr}). Для другой станции скачайте отдельный шаблон.",
            "Заполните первый лист: каждая строка — одна запись. Текущие данные уже внесены, их можно исправлять.",
            "В ячейках со списком выбирайте значение из списка. Даты: ДД.ММ.ГГГГ.",
            "При загрузке файла данные этого раздела станции полностью заменяются данными из файла.",
            "Если есть ошибки, система покажет строку и столбец и ничего не изменит.",
        ],
    },
    "en": {
        "roads": "Access roads", "rleader": "Manager", "leader": "Full name", "rlength": "Length (m)", "contract": "Contract", "cargo": "Cargo operation time (h)",
        "owner": "Road owner", "inn": "TIN", "email": "E-mail", "distance_km": "Delivery distance (km)",
        "length_t1": "T1", "length_t2": "T2", "contract_date": "Signed", "contract_start": "Start", "contract_end": "End",
        "has_tech_passport": "Tech passport", "has_instruction": "Instruction", "junction": "Junction", "load_hours": "Loading",
        "unload_hours": "Unloading", "oper2_hours": "Double operation", "front": "Front (wagons)", "wagons_at_once": "Wagons per delivery",
        "created_year": "Year created", "canceled": "Terminated",
        "tracks": "Parks and tracks", "workers": "Staff list", "hint": "* marks required columns. Enter data from row 5. Do not reorder columns.",
        "arrows": "Limiting switches", "length": "Length (m)", "devices": "Equipment",
        "park": "Park", "track_no": "Track №", "purpose": "Name and purpose", "arrow_from": "from", "arrow_to": "to",
        "length_total": "Total", "usable_length": "Usable", "wagon_capacity": "Wagon capacity", "insulation": "Insulation",
        "contact_network": "Contact network", "alns": "ALSN",
        "full_name": "Full name", "position": "Position", "gender": "Gender", "birth_date": "Birth date", "education": "Education",
        "contract_type": "Contract type", "hired_date": "Hired on", "phone": "Phone",
        "yes": "Yes", "no": "No", "M": "Male", "F": "Female", "higher": "Higher", "second": "Second degree",
        "special": "Vocational", "secondary": "Secondary",
        "guide": "Instructions", "g": [
            "This template is for station {station} (ESR {esr}). Download a separate template for another station.",
            "Fill in the first sheet: one row per record. Current data is already filled in and can be edited.",
            "Pick values from the list in list cells. Dates: DD.MM.YYYY.",
            "Uploading the file fully replaces this section's data for the station with the file's rows.",
            "If there are errors, the system shows the row and column and changes nothing.",
        ],
    },
}
PARKS = ["Jo'natish parki", "Qabul qilish parki", "Saralash parki", "Asosiy park"]
CONTRACTS = ["Asosiy", "Ortiqcha ish o'rni", "O'rindoshlik", "Muddatli"]
ROAD_CONTRACTS = ["Qo'yib olish", "Tashish", "Ijaraga", "Boshqa"]

HEAD = PatternFill("solid", fgColor="0E7490")
GROUP = PatternFill("solid", fgColor="0B2530")
REQ = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="CBD5E1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _label(lang: str, kind: str, value) -> str:
    """Display text for a stored value in the template language."""
    L = T[lang]
    if kind == "bool":
        return L["yes"] if value else L["no"]
    if kind == "choice" and isinstance(value, str):
        return L.get(value, value)
    return value


def build_template(kind: str, rows: list[dict], station: str, esr: str, lang: str = "uz") -> bytes:
    lang = lang if lang in T else "uz"
    L, spec = T[lang], SPECS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = L[kind][:31]
    n = len(spec)
    last = get_column_letter(n)

    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = f"SilkRail · {station} (ESR {esr}) · {L[kind]}"
    ws["A1"].font = Font(bold=True, size=14, color="0B2530")
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = L["hint"]
    ws["A2"].font = Font(italic=True, size=10, color="64748B")

    # two header rows: group titles (row 3) over column titles (row 4), like the table on the site
    col = 1
    while col <= n:
        key, width, ck, req, _, group = spec[col - 1]
        if group:
            end = col
            while end < n and spec[end][5] == group:
                end += 1
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=end)
            c = ws.cell(3, col, L[group])
            for k in range(col, end + 1):
                ws.cell(3, k).fill = GROUP
                ws.cell(3, k).border = BOX
            c.font, c.alignment = Font(bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="center")
            for k in range(col, end + 1):
                sk = spec[k - 1]
                h = ws.cell(4, k, L[sk[0]] + (" *" if sk[3] else ""))
                h.fill, h.border = HEAD, BOX
                h.font, h.alignment = Font(bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="center", wrap_text=True)
            col = end + 1
        else:
            ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
            h = ws.cell(3, col, L[key] + (" *" if req else ""))
            ws.cell(4, col).border = BOX
            h.fill, h.border = HEAD, BOX
            h.font, h.alignment = Font(bold=True, color="FFFFFF"), Alignment(horizontal="center", vertical="center", wrap_text=True)
            col += 1
    for i, (key, width, *_rest) in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = f"A{FIRST_ROW}"

    # current data
    for r, row in enumerate(rows, start=FIRST_ROW):
        for i, (key, _, ck, *_rest) in enumerate(spec, start=1):
            c = ws.cell(r, i, _label(lang, ck, row.get(key)))
            if ck == "date":
                c.number_format = "DD.MM.YYYY"
    end_row = FIRST_ROW + max(len(rows), 0) + BLANK_ROWS
    for r in range(FIRST_ROW, end_row + 1):
        for i, (_, _, ck, req, *_rest) in enumerate(spec, start=1):
            c = ws.cell(r, i)
            c.border = BOX
            if req:
                c.fill = REQ
            if ck == "date":
                c.number_format = "DD.MM.YYYY"

    # validations
    def rng(i):
        L_ = get_column_letter(i)
        return f"{L_}{FIRST_ROW}:{L_}{end_row}"

    def dv_list(values, i, strict=True):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=strict, errorStyle="stop" if strict else "information")
        ws.add_data_validation(dv)
        dv.add(rng(i))

    for i, (key, _, ck, req, choices, _g) in enumerate(spec, start=1):
        if ck == "bool":
            dv_list([L["yes"], L["no"]], i)
        elif ck == "choice":
            dv_list([L.get(c, c) for c in choices], i)
        elif ck == "int":
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True)
            ws.add_data_validation(dv)
            dv.add(rng(i))
        elif ck == "date":
            dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True, showErrorMessage=True)
            ws.add_data_validation(dv)
            dv.add(rng(i))
        elif key == "park":
            dv_list(PARKS, i, strict=False)  # suggestions; other park names are allowed
        elif key == "contract_type":
            dv_list(ROAD_CONTRACTS if kind == "roads" else CONTRACTS, i, strict=False)
        elif ck == "float":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True)
            ws.add_data_validation(dv)
            dv.add(rng(i))

    g = wb.create_sheet(L["guide"][:31])
    g.column_dimensions["A"].width = 110
    g["A1"] = L["guide"]
    g["A1"].font = Font(bold=True, size=13)
    for k, line in enumerate(L["g"], start=3):
        g.cell(k, 1, f"{k - 2}. " + line.format(station=station, esr=esr)).alignment = Alignment(wrap_text=True)

    meta = wb.create_sheet("_meta")
    meta["A1"], meta["B1"] = "silkrail-template", kind
    meta["A2"], meta["B2"] = "esr", esr
    meta["A3"], meta["B3"] = "version", 1
    meta.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TemplateError(Exception):
    def __init__(self, code: str, **info):
        super().__init__(code)
        self.code, self.info = code, info


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError


def parse_template(kind: str, content: bytes, esr: str) -> list[dict]:
    """Returns clean rows or raises TemplateError("invalid", errors=[...]) without touching anything."""
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise TemplateError("bad_file")
    if "_meta" not in wb.sheetnames or wb["_meta"]["A1"].value != "silkrail-template":
        raise TemplateError("not_template")
    meta = wb["_meta"]
    if meta["B1"].value != kind:
        raise TemplateError("wrong_kind")
    if str(meta["B2"].value) != str(esr):
        raise TemplateError("wrong_station", fileEsr=str(meta["B2"].value))
    ws = wb.worksheets[0]
    spec = SPECS[kind]
    out, errors = [], []

    def err(r, i, code, value=None):
        errors.append({"row": r, "col": get_column_letter(i), "field": spec[i - 1][0], "code": code,
                       "value": None if value is None else str(value)[:60]})

    for r in range(FIRST_ROW, min(ws.max_row, FIRST_ROW + MAX_ROWS) + 1):
        raw = [ws.cell(r, i).value for i in range(1, len(spec) + 1)]
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        row = {}
        for i, ((key, _, ck, req, choices, _g), v) in enumerate(zip(spec, raw), start=1):
            empty = v is None or str(v).strip() == ""
            if empty:
                if req:
                    err(r, i, "required")
                row[key] = False if ck == "bool" else None
                continue
            try:
                if ck == "int":
                    f = float(str(v).replace(" ", "").replace(",", "."))
                    if f < 0 or f != int(f):
                        raise ValueError
                    row[key] = int(f)
                elif ck == "float":
                    f = float(str(v).replace(" ", "").replace(",", "."))
                    if f < 0:
                        raise ValueError
                    row[key] = round(f, 2)
                elif ck == "date":
                    d = _parse_date(v)
                    # contracts may legitimately end decades ahead; a birth date cannot be in the future
                    if not (1900 <= d.year <= date.today().year + 50) or (key == "birth_date" and d > date.today()):
                        raise ValueError
                    row[key] = d
                elif ck == "bool":
                    nv = _norm(v)
                    if nv in BOOL[True]:
                        row[key] = True
                    elif nv in BOOL[False]:
                        row[key] = False
                    else:
                        raise ValueError
                elif ck == "choice":
                    nv = _norm(v)
                    labels = {_norm(T[lg].get(c, c)): c for lg in T for c in choices}
                    hit = labels.get(nv) or next((c for c, al in choices.items() if nv in al or nv == _norm(c)), None)
                    if not hit:
                        raise ValueError
                    row[key] = hit
                else:
                    row[key] = str(v).strip()[:160]
            except (ValueError, TypeError):
                err(r, i, {"int": "number", "float": "decimal", "date": "date"}.get(ck, "choice"), v)
        if kind == "tracks" and isinstance(row.get("length_total"), int) and isinstance(row.get("usable_length"), int) \
                and row["usable_length"] > row["length_total"]:
            err(r, 7, "usable_gt_total", row["usable_length"])
        if kind == "roads":
            col = {k[0]: n for n, k in enumerate(spec, start=1)}
            lt, t1, t2 = row.get("length_total"), row.get("length_t1") or 0, row.get("length_t2") or 0
            if isinstance(lt, float) and t1 + t2 > lt + 0.5:
                err(r, col["length_t1"], "parts_gt_total", f"{t1} + {t2}")
            cs, ce = row.get("contract_start"), row.get("contract_end")
            if isinstance(cs, date) and isinstance(ce, date) and ce < cs:
                err(r, col["contract_end"], "end_before_start", ce)
            y = row.get("created_year")
            if isinstance(y, int) and not 1900 <= y <= date.today().year:
                err(r, col["created_year"], "year", y)
        if kind == "workers" and isinstance(row.get("birth_date"), date) and isinstance(row.get("hired_date"), date) \
                and row["hired_date"] < row["birth_date"]:
            err(r, 7, "hired_before_birth", row["hired_date"])
        out.append(row)
        if len(errors) >= 100:
            break
    if errors:
        raise TemplateError("invalid", errors=errors)
    if not out:
        raise TemplateError("empty")
    return out
